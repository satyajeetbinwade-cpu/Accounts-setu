"""Deterministic parse of a books-side Credit/Debit NOTE register.

The fixture reproduces the reference workbook's shape EXACTLY as described by
the user (two sheets, 20 columns, the five narration spellings, sparse rate
buckets, a sign-inconsistent `Short & Excess`) — but with anonymised
suppliers, so no real client GSTIN/PAN/name is committed.

Run:

    venv/bin/python -m pytest tests/test_note_register.py -v
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from src.data_paths import VALID_SOURCE_TYPES
from src.f6.note_register import (
    NOTE_REQUIRED_FIELDS,
    NOTE_SHEET_KEY,
    NOTE_SOURCE_TYPES,
    _ALIASES,
    _check_bucket_explained,
    parse_note_register,
)
from src.f6.rate_buckets import RateBucketMatrix, detect_rate_buckets
from src.shared import discovery

NOTE_SHEET = "Debit Note Register"
BOOKS_SHEET = "Purchase Register"

# Column order verbatim from the real sheet.
HEADERS = [
    "Date", "Particulars", "Buyer", "Buyer Address", "Voucher Type",
    "Voucher No.", "Voucher Ref. No.", "Voucher Ref. Date", "GSTIN/UIN",
    "Sales Tax No.", "PAN No.", "Narration", "Gross Total",
    "Pur- Lehanga", "Pur - Sarees",
    "Unclaimed CGST @ 2.5%", "Unclaimed SGST @ 2.5%",
    "Unclaimed CGST @ 6%", "Unclaimed SGST @ 6%", "Short & Excess",
]

# The four representative rows, with the formatting quirks preserved:
# "Voucher Ref. No." blank on EVERY row; "Voucher Type" = Debit Note on every
# row while every narration cites a supplier credit note; ISO dates.
ROWS = [
    ["2025-04-09", "Supplier A", "Supplier A", "Addr A", "Debit Note", "8/2025-26",
     None, "2025-03-26", "07AAAAA0000A1ZZ", None, "AAAAA0000A",
     "CREDIT NOTE NO.  CD-12/04-2025", 5235, 2185.38, 2654.64,
     66.37, 66.37, 131.12, 131.12, None],
    ["2025-04-22", "Supplier B", "Supplier B", "Addr B", "Debit Note", "3/2025-26",
     None, "2025-03-26", "07BBBBB0000B1ZZ", None, "BBBBB0000B",
     "CREDIT NOTE NO. SRI-312/25-26", 22583, 7995, 12980,
     324.51, 324.51, 479.7, 479.7, 0.42],
    ["2025-04-12", "Supplier C", "Supplier C", "Addr C", "Debit Note", "7/2025-26",
     None, "2025-03-25", "07CCCCC0000C1ZZ", None, "CCCCC0000C",
     "CREDIT NOTE NO. 50", 4711, 4206, None,
     None, None, 252.36, 252.36, 0.28],
    ["2025-04-23", "Supplier A", "Supplier A", "Addr A", "Debit Note", "2/2025-26",
     None, "2025-03-26", "07AAAAA0000A1ZZ", None, "AAAAA0000A",
     "Credit Note No CD-197/04-2025", 5125, 4575.9, None,
     None, None, 274.55, 274.55, None],
]


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    path = tmp_path_factory.mktemp("note_register") / "Credit Note.xlsx"
    book = Workbook()

    # Sheet 1 — the EXISTING Purchase Register input. A different layout
    # entirely; parsing it as notes would be obvious.
    books = book.active
    books.title = BOOKS_SHEET
    books.append(["Bill No", "Date", "Party", "GSTIN", "Taxable",
                  "CGST Tax Amt @ 9%", "SGST Tax Amt @ 9%", "Bill Amount"])
    books.append(["B/1", "2025-04-05", "Supplier X", "07XXXXX0000X1ZZ", 1000, 90, 90, 1180])

    # Sheet 2 — the note register, behind a one-row title block so header
    # detection is genuinely exercised.
    notes = book.create_sheet(NOTE_SHEET)
    notes.append(["Debit Note Register"])
    notes.append(HEADERS)
    for row in ROWS:
        notes.append(row)

    book.save(path)
    return path


@pytest.fixture(scope="module")
def parsed(workbook):
    result = parse_note_register(str(workbook), "credit_notes", "Credit Note.xlsx")
    assert result is not None, "the reference shape must parse deterministically"
    return result


# --- Slot registration ------------------------------------------------------


def test_credit_notes_slot_is_registered():
    assert "credit_notes" in VALID_SOURCE_TYPES
    assert "credit_notes" in NOTE_SOURCE_TYPES
    assert discovery.SOURCE_TYPE_LABELS["credit_notes"] == "Credit / Debit Note Register (received)"
    assert discovery.source_type_hint("credit_notes")


def test_note_slot_is_never_a_recon_slot():
    """The note register must never be picked as the books or portal side of
    an INVOICE run — it is a separate document class."""
    for recon_type, sources in discovery.RECON_SOURCE_TYPES.items():
        assert "credit_notes" not in sources, recon_type


def test_aliases_agree_with_the_documented_config_contract():
    """The `gst.notes.columns` block is the documented contract; the module's
    alias sets are the implementation. They must not drift."""
    from src.config_loader import load_config

    columns = load_config()["gst"]["notes"]["columns"]
    value_block = load_config()["gst"]["notes"]["value_block"]
    assert columns["narration_aliases"] == _ALIASES["narration"]
    assert columns["party_aliases"] == _ALIASES["party_name"]
    assert columns["gstin_aliases"] == _ALIASES["gstin"]
    assert columns["reference_aliases"] == _ALIASES["voucher_ref"]
    assert columns["note_date_aliases"] == _ALIASES["invoice_date"]
    assert columns["booking_date_aliases"] == _ALIASES["booking_date"]
    assert columns["voucher_type_aliases"] == _ALIASES["voucher_type"]
    assert columns["voucher_number_aliases"] == _ALIASES["voucher_no"]
    assert value_block["invoice_value_column_aliases"] == _ALIASES["invoice_value"]
    assert value_block["rounding_column_aliases"] == _ALIASES["rounding_adjustment"]
    # "Particulars" is the SUPPLIER NAME in this layout, never the narration.
    assert "particulars" not in _ALIASES["narration"]
    assert "particulars" in _ALIASES["party_name"]


# --- Rate-bucket detector extension ----------------------------------------


def test_loose_bucket_headers_detected_as_tax():
    """`Unclaimed CGST @ 2.5%` has a prefix and NO kind token, so the strict
    shape cannot match it. Arithmetic proved these are TAX columns."""
    matrix = detect_rate_buckets([
        "Unclaimed CGST @ 2.5%", "Unclaimed SGST @ 2.5%",
        "Unclaimed CGST @ 6%", "Unclaimed SGST @ 6%",
    ])
    assert matrix.detected
    assert [(b.head, b.kind, b.rate, b.prefix) for b in matrix.buckets] == [
        ("CGST", "tax", 2.5, "unclaimed"),
        ("SGST", "tax", 2.5, "unclaimed"),
        ("CGST", "tax", 6.0, "unclaimed"),
        ("SGST", "tax", 6.0, "unclaimed"),
    ]
    assert matrix.has_taxable_buckets() is False
    assert len(matrix.tax_mirror_pairs()) == 2
    # CGST is the mirror's source of truth for the explanation check.
    assert [b["rate"] for b in matrix.explanation_inputs()] == [2.5, 6.0]


def test_purchase_register_bucket_shape_is_unchanged():
    """The strict shape must still win, byte for byte, so the Purchase
    Register path cannot change meaning."""
    matrix = detect_rate_buckets([
        "CGST Txbl. Amt @ 2.5%", "CGST Tax Amt @ 2.5%",
        "SGST Txbl. Amt @ 2.5%", "SGST Tax Amt @ 2.5%",
    ])
    assert [(b.head, b.kind, b.rate) for b in matrix.buckets] == [
        ("CGST", "taxable", 2.5), ("CGST", "tax", 2.5),
        ("SGST", "taxable", 2.5), ("SGST", "tax", 2.5),
    ]
    assert all(b.prefix == "" for b in matrix.buckets), "strict matches carry no prefix"
    assert matrix.has_taxable_buckets() is True
    # §5.4's taxable de-duplication is untouched.
    rules = {r["canonical_field"]: r for r in matrix.aggregate_rules()}
    assert rules["taxable_value"]["excluded_mirrors"] == ["SGST Txbl. Amt @ 2.5%"]


def test_a_plain_column_is_never_a_bucket():
    matrix = detect_rate_buckets(["cgst", "sgst", "igst", "Taxable Value", "Gross Total"])
    assert not matrix.detected


# --- Parsing the reference shape -------------------------------------------


def test_reads_the_note_sheet_not_the_purchase_register(parsed):
    """One workbook carries BOTH registers — the note parser must take the
    note sheet and leave the books sheet alone."""
    assert len(parsed.rows) == 4
    refs = {str(row["note_reference"]) for row in parsed.rows}
    assert refs == {"CD12042025", "SRI312", "50", "CD197042025"}


def test_item_category_columns_detected_structurally(parsed):
    """Detected by POSITION, not by name — the categories are garment names
    with inconsistent spacing, so a name-based rule would break the next
    client."""
    assert parsed.item_category_columns == ["Pur- Lehanga", "Pur - Sarees"]


def test_taxable_value_is_the_sum_of_item_categories(parsed):
    frame = parsed.frame.set_index("note_reference")
    assert float(frame.loc["CD12042025", "taxable_value"]) == pytest.approx(4840.02)
    assert float(frame.loc["SRI312", "taxable_value"]) == pytest.approx(20975.00)
    assert float(frame.loc["50", "taxable_value"]) == pytest.approx(4206.00)
    assert float(frame.loc["CD197042025", "taxable_value"]) == pytest.approx(4575.90)


def test_rate_buckets_sum_into_the_tax_heads(parsed):
    frame = parsed.frame.set_index("note_reference")
    assert float(frame.loc["CD12042025", "cgst"]) == pytest.approx(197.49)  # 66.37 + 131.12
    assert float(frame.loc["CD12042025", "sgst"]) == pytest.approx(197.49)
    assert float(frame.loc["CD12042025", "total_tax"]) == pytest.approx(394.98)
    assert float(frame.loc["CD12042025", "invoice_value"]) == pytest.approx(5235.00)


def test_igst_is_absent_not_zero(parsed):
    """An intra-state register has no IGST column. NULL, never 0 — 0 would
    read as 'zero IGST charged'."""
    assert all(value is None for value in parsed.frame["igst"])


def test_every_row_is_a_credit_note_despite_the_voucher_label(parsed):
    """The whole point: 100% Credit Notes, zero genuine debit notes, even
    though EVERY row's voucher-type label says "Debit Note"."""
    assert {row["note_kind"] for row in parsed.rows} == {"Credit Note"}
    assert {row["voucher_type_raw"] for row in parsed.rows} == {"Debit Note"}
    assert all(row["note_reference_signal"] == "citation" for row in parsed.rows)


def test_the_booking_voucher_number_is_never_the_reference(parsed):
    """`Voucher No.` is the client's own booking number. Mapping it to the
    join key would fail SILENTLY — every note simply unmatched."""
    refs = {str(row["note_reference"]) for row in parsed.rows}
    assert "8202526" not in refs
    assert not any(v in "8/2025-26" for v in refs)


def test_note_date_is_the_reference_date_not_the_booking_date(parsed):
    """The note's own date is a period BEHIND the booking, so the booking
    Date must not be used for the date-window fallback."""
    frame = parsed.frame.set_index("note_reference")
    assert str(frame.loc["CD12042025", "note_date"]) == "2025-03-26"
    assert str(frame.loc["CD12042025", "note_date"]) != "2025-04-09"
    assert parsed.metadata.date_format == "%Y-%m-%d"


def test_rounding_residual_is_mapped_and_both_directions_survive(parsed):
    """`Short & Excess` is over on one row and short on the next; both must
    parse without a hard stop."""
    frame = parsed.frame.set_index("note_reference")
    assert float(frame.loc["SRI312", "rounding_adjustment"]) == pytest.approx(0.42)
    assert float(frame.loc["50", "rounding_adjustment"]) == pytest.approx(0.28)


def test_no_row_is_dropped_and_no_row_is_a_duplicate(parsed):
    """Rows 1 and 4 share a GSTIN AND a note date but are different notes —
    the duplicate key must be (GSTIN, reference), not the invoice key."""
    assert len(parsed.frame) == len(parsed.rows) == 4


def test_columns_are_all_dispositioned(parsed):
    dispositions = {d["column"]: d for d in parsed.column_dispositions}
    assert dispositions["Voucher No."]["disposition"] == "validation-signal"
    assert dispositions["Date"]["disposition"] == "validation-signal"
    assert dispositions["Pur- Lehanga"]["canonical_field"] == "taxable_value"
    assert dispositions["Unclaimed CGST @ 2.5%"]["canonical_field"] == "cgst"
    assert not [d for d in parsed.column_dispositions if d["disposition"] == "needs-decision"]


# --- Validation -------------------------------------------------------------


def test_validation_passes_on_the_real_shape(parsed):
    outcome = parsed.run_validation(required_fields=NOTE_REQUIRED_FIELDS)
    by_check = {result.check: result for result in outcome.results}
    assert by_check["rate_bucket_explained"].result == "pass"
    assert by_check["mirror_columns"].result == "pass"
    assert by_check["tax_arithmetic"].result == "pass"
    assert by_check["required_fields"].result == "pass"
    assert by_check["duplicate_detection"].result == "pass"
    assert by_check["row_accounting"].result == "pass"
    assert not outcome.hard_stopped


def test_igst_is_not_required(parsed):
    """Requiredness must not demand a column this layout never has."""
    assert "igst" not in NOTE_REQUIRED_FIELDS


def test_explanation_check_passes_a_bucket_its_item_column_explains():
    rows = [{"invoice_value": 1180, "Pur": 1000, "Unclaimed CGST @ 9%": 90}]
    result = _check_bucket_explained(
        rows, ["Pur"], [{"rate": 9.0, "tax_column": "Unclaimed CGST @ 9%", "head": "CGST"}]
    )
    assert result.result == "pass"


def test_explanation_check_accepts_a_SUBSET_of_item_columns():
    """Two categories can share one rate, so the bucket covers their SUM —
    requiring a single column would false-flag those rows."""
    rows = [{"invoice_value": 2240, "Cat A": 1000, "Cat B": 1000,
             "Unclaimed CGST @ 6%": 120}]
    result = _check_bucket_explained(
        rows, ["Cat A", "Cat B"],
        [{"rate": 6.0, "tax_column": "Unclaimed CGST @ 6%", "head": "CGST"}],
    )
    assert result.result == "pass"


def test_explanation_check_flags_a_bucket_no_item_column_explains():
    """A bucket filed under the wrong rate (or a taxable column mistaken for a
    tax column) is exactly what this replaces the implied-rate check for."""
    rows = [{"invoice_value": 1180, "Pur": 1000, "Unclaimed CGST @ 18%": 90}]
    result = _check_bucket_explained(
        rows, ["Pur"], [{"rate": 18.0, "tax_column": "Unclaimed CGST @ 18%", "head": "CGST"}]
    )
    assert result.result == "hard_stop", "1 of 1 rows fails -> escalates"


def test_explanation_check_row_flags_below_the_escalation_threshold():
    rows = [{"invoice_value": 10, "Pur": 1000.0, "Unclaimed CGST @ 9%": 90.0} for _ in range(5)]
    rows[0]["Unclaimed CGST @ 9%"] = 1.0  # 1 of 5 rows (20%) — not above the threshold
    result = _check_bucket_explained(
        rows, ["Pur"], [{"rate": 9.0, "tax_column": "Unclaimed CGST @ 9%", "head": "CGST"}]
    )
    assert result.result == "row_flag"


# --- Fall-through -----------------------------------------------------------


def test_a_books_workbook_with_no_note_sheet_falls_through(tmp_path):
    """Unknown shape ⇒ None, so the caller reaches the generic AI-proposal +
    human-confirmation path instead of force-fitting."""
    path = tmp_path / "books.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = BOOKS_SHEET
    sheet.append(["Bill No", "Date", "Party", "GSTIN", "Bill Amount"])
    sheet.append(["B/1", "2025-04-05", "Supplier X", "07XXXXX0000X1ZZ", 1180])
    book.save(path)
    assert parse_note_register(str(path), "credit_notes", "books.xlsx") is None


def test_a_note_sheet_with_no_rate_buckets_falls_through(tmp_path):
    path = tmp_path / "flat.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = NOTE_SHEET
    sheet.append(["Date", "Narration", "Gross Total", "Pur- Lehanga"])
    sheet.append(["2025-04-09", "CREDIT NOTE NO. CD-1", 100, 100])
    book.save(path)
    assert parse_note_register(str(path), "credit_notes", "flat.xlsx") is None


def test_other_source_types_are_not_parsed_as_notes(workbook):
    assert parse_note_register(str(workbook), "tally", "Credit Note.xlsx") is None


# --- The ingestion entry point ---------------------------------------------


def test_normalize_source_file_routes_credit_notes_to_the_note_parser(workbook, tmp_path):
    """The slot must reach the deterministic parser through the ONE ingestion
    entry point, with ZERO model calls."""
    from src import bootstrap
    from src.ingestion_ai import normalizer

    db = tmp_path / "note_smoke.db"
    bootstrap.init_all(str(db))

    result = normalizer.normalize_source_file(
        str(workbook), "credit_notes", "Test Client", "2025-04",
        actor="test", filename="Credit Note.xlsx", db_path=str(db),
    )

    assert result.ingestion_path == "deterministic"
    assert result.model_used is None
    assert result.llm_cached is False
    assert not result.hard_stopped
    assert result.status in ("ok", "partial")
    assert len(result.canonical_df) == 4
    assert "note_reference" in result.canonical_df.columns
    assert "note_kind" in result.canonical_df.columns
    assert any("SEPARATE document class" in note for note in result.notes)
    # A note register has no invoice number, and an intra-state one has no
    # IGST — neither may be reported as an unmapped REQUIRED field.
    assert set(result.unmapped_required_fields) == set()


def test_normalize_source_file_leaves_the_note_register_alone_for_other_slots(workbook, tmp_path):
    """The SAME workbook in the books slot must not be parsed as notes — the
    note parser is gated on the slot, not on the file looking note-ish."""
    from src import bootstrap
    from src.ingestion_ai import normalizer

    db = tmp_path / "books_smoke.db"
    bootstrap.init_all(str(db))

    result = normalizer.normalize_source_file(
        str(workbook), "tally", "Test Client", "2025-04",
        actor="test", filename="Credit Note.xlsx", db_path=str(db),
    )
    assert result.source_type == "tally"
    # The note parser must not have run: no note-identity columns, and none of
    # its explanatory notes.
    assert "note_kind" not in result.canonical_df.columns
    assert "note_reference" not in result.canonical_df.columns
    assert not any("NOTE-register" in note for note in result.notes)


# --- The shared frame builder must not change for the books path -----------


def test_to_frame_default_identity_filter_is_unchanged():
    from src.f6.books_register import _to_frame

    frame = _to_frame(
        [{"gstin": "", "invoice_number": ""}, {"gstin": "07X", "invoice_number": "A1"}],
        "tally", "f.xlsx", RateBucketMatrix(),
    )
    assert len(frame) == 1
    assert frame.iloc[0]["invoice_number"] == "A1"
