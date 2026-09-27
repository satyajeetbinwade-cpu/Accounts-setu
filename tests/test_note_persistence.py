"""Note results persist in their OWN table, and the note pass runs only when a
books register was supplied.

The structural claim this file exists to prove: a note can NEVER appear in the
invoice results, because it is never written to `match_results` at all.

Run:

    venv/bin/python -m pytest tests/test_note_persistence.py -v
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from src import bootstrap, queries
from src import db as recon_db
from src.config_loader import load_config

CLIENT = "Note Client"
PERIOD = "2025-04"

HEADERS = [
    "Date", "Particulars", "Buyer", "Buyer Address", "Voucher Type",
    "Voucher No.", "Voucher Ref. No.", "Voucher Ref. Date", "GSTIN/UIN",
    "Sales Tax No.", "PAN No.", "Narration", "Gross Total",
    "Pur- Lehanga", "Pur - Sarees",
    "Unclaimed CGST @ 2.5%", "Unclaimed SGST @ 2.5%",
    "Unclaimed CGST @ 6%", "Unclaimed SGST @ 6%", "Short & Excess",
]

ROWS = [
    ["2025-04-09", "Supplier A", "Supplier A", "Addr A", "Debit Note", "8/2025-26",
     None, "2025-03-26", "07AAAAA0000A1ZZ", None, "AAAAA0000A",
     "CREDIT NOTE NO.  CD-12/04-2025", 5235, 2185.38, 2654.64,
     66.37, 66.37, 131.12, 131.12, None],
    ["2025-04-22", "Supplier B", "Supplier B", "Addr B", "Debit Note", "3/2025-26",
     None, "2025-03-26", "07BBBBB0000B1ZZ", None, "BBBBB0000B",
     "CREDIT NOTE NO. SRI-312/25-26", 22583, 7995, 12980,
     324.51, 324.51, 479.7, 479.7, 0.42],
]


@pytest.fixture
def note_register(tmp_path):
    """A note register filed under a private DATA_ROOT."""
    root = tmp_path / "data"
    folder = root / CLIENT / PERIOD / "credit_notes"
    folder.mkdir(parents=True, exist_ok=True)

    path = folder / "Credit Note.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Debit Note Register"
    sheet.append(["Debit Note Register"])
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)
    book.save(path)

    import src.data_paths as data_paths

    original = data_paths.DATA_ROOT
    data_paths.DATA_ROOT = root
    try:
        yield {"credit_notes": path.name}
    finally:
        data_paths.DATA_ROOT = original


@pytest.fixture
def db(tmp_path):
    path = tmp_path / "notes.db"
    bootstrap.init_all(str(path))
    return path


# --- Storage separation ----------------------------------------------------


def test_note_results_live_in_their_own_table(db):
    """The invoice surface reads `match_results`; a note is never written
    there — so no filter can be the thing that protects the invoice counts."""
    conn = recon_db.get_connection(str(db))
    try:
        run_id = recon_db.insert_run(
            conn, client=CLIENT, period=PERIOD, recon_type="GST",
            run_timestamp="2025-05-01T00:00:00Z", config_snapshot="{}",
            source_file_names=["Credit Note.xlsx"],
        )
        recon_db.insert_note_results(conn, [{
            "run_id": run_id, "classification": "Not in Portal", "confidence_score": 0.0,
            "confidence_band": "Low", "note_reference": "CD12042025",
            "note_reference_raw": "CD-12/04-2025", "note_kind": "Credit Note",
            "note_date": "2025-03-26", "note_tax": 394.98, "note_value": 5235.0,
            "note_side": "books", "difference_type": None,
            "books_record": '{"gstin": "07AAAAA0000A1ZZ"}', "portal_record": None,
            "match_reason": "Not in Portal: no counterpart.",
        }])
        conn.commit()
    finally:
        conn.close()

    notes = queries.get_note_results(run_id, db_path=str(db))
    assert len(notes) == 1
    assert notes.iloc[0]["note_reference"] == "CD12042025"
    assert notes.iloc[0]["note_tax"] == pytest.approx(394.98)
    assert notes.iloc[0]["books_record"] == {"gstin": "07AAAAA0000A1ZZ"}

    # THE POINT: the invoice results are untouched by a note.
    assert len(queries.get_results(run_id, db_path=str(db))) == 0
    assert queries.get_run_summary(run_id, db_path=str(db))["total_results"] == 0


def test_note_results_are_empty_not_missing_for_a_run_without_notes(db):
    conn = recon_db.get_connection(str(db))
    try:
        run_id = recon_db.insert_run(
            conn, client=CLIENT, period=PERIOD, recon_type="GST",
            run_timestamp="2025-05-01T00:00:00Z", config_snapshot="{}",
            source_file_names=[],
        )
        recon_db.insert_note_results(conn, [])
        conn.commit()
    finally:
        conn.close()
    assert len(queries.get_note_results(run_id, db_path=str(db))) == 0


# --- The pass inside a run -------------------------------------------------


def test_the_note_pass_runs_when_a_register_is_supplied(note_register, db):
    from src.runner import _run_note_pass

    run_notes: list[dict] = []
    results = _run_note_pass(
        CLIENT, PERIOD, load_config()["gst"], note_register, run_notes,
        client_id=None, actor="test", db_path=str(db),
    )

    assert len(results) == 2, "both register rows must become note results"
    assert {r["classification"] for r in results} == {"Not in Portal"}
    assert {r["note_kind"] for r in results} == {"Credit Note"}
    assert {r["note_reference"] for r in results} == {"CD12042025", "SRI312"}

    declared = [n for n in run_notes if n["code"] == "note_reconciliation"]
    assert len(declared) == 1
    assert "2 note(s) were reconciled separately from the invoices" in declared[0]["title"]
    assert "never enter the invoice table" in declared[0]["detail"]


def test_the_note_pass_is_skipped_when_no_register_was_supplied(db):
    """With no register the honest output is the caveat, NOT a wall of false
    Not-in-Books rows for every portal note."""
    from src.runner import _run_note_pass

    run_notes: list[dict] = []
    results = _run_note_pass(
        CLIENT, PERIOD, load_config()["gst"], {}, run_notes,
        client_id=None, actor="test", db_path=str(db),
    )
    assert results == []
    assert run_notes == []


def test_a_missing_register_file_degrades_to_no_pass(db, tmp_path):
    """A register that is selected but not on disk must not crash the run."""
    from src.runner import _run_note_pass

    run_notes: list[dict] = []
    results = _run_note_pass(
        CLIENT, PERIOD, load_config()["gst"], {"credit_notes": "nope.xlsx"}, run_notes,
        client_id=None, actor="test", db_path=str(db),
    )
    assert results == []


def test_probing_for_a_previous_period_does_not_create_a_folder(note_register, db, tmp_path):
    """`load_note_frames` looks back a period for late notes — that probe must
    not leave an empty period folder behind (it would appear on every picker)."""
    from src.ingestion_ai.normalizer import load_note_frames
    from src.data_paths import DATA_ROOT

    load_note_frames(CLIENT, PERIOD, selected_files=note_register, actor="test", db_path=str(db))
    assert not (DATA_ROOT / CLIENT / "2025-03").exists()


def test_prev_period_walks_back_across_a_year_boundary():
    from src.ingestion_ai.normalizer import _prev_period

    assert _prev_period("2025-04") == "2025-03"
    assert _prev_period("2025-01") == "2024-12"
    assert _prev_period("-") is None
    assert _prev_period(None) is None
    assert _prev_period("2025-13") == "2025-12"  # tolerated, never crashes


# --- Report plumbing -------------------------------------------------------


def test_the_register_frame_carries_the_note_columns_end_to_end(note_register, db):
    from src.ingestion_ai.normalizer import load_note_frames

    books, portal, adjacent, warnings = load_note_frames(
        CLIENT, PERIOD, selected_files=note_register, actor="test", db_path=str(db),
    )
    assert warnings == []
    assert portal is None and adjacent is None
    assert books is not None
    for column in ("note_reference", "note_reference_raw", "note_kind", "note_date"):
        assert column in books.columns, column
    assert books.iloc[0]["note_kind"] == "Credit Note"
    assert books.iloc[0]["note_reference"] == "CD12042025"


# --- The report reads the NOTE PASS, not the portal file -------------------


@pytest.fixture
def run_with_notes(db):
    """A run carrying real note results and NO invoice results."""
    conn = recon_db.get_connection(str(db))
    try:
        run_id = recon_db.insert_run(
            conn, client=CLIENT, period=PERIOD, recon_type="GST",
            run_timestamp="2025-05-01T00:00:00Z", config_snapshot="{}",
            source_file_names=["Credit Note.xlsx"],
        )
        recon_db.insert_note_results(conn, [
            {
                "run_id": run_id, "classification": "Not in Portal",
                "confidence_score": 0.0, "confidence_band": "Low",
                "note_reference": "CD12042025", "note_reference_raw": "CD-12/04-2025",
                "note_kind": "Credit Note", "note_date": "2025-03-26",
                "note_tax": 394.98, "note_value": 5235.0, "note_side": "books",
                "difference_type": None,
                "books_record": '{"gstin": "07AAAAA0000A1ZZ", "party_name": "Supplier A",'
                                ' "taxable_value": 4840.02, "cgst": 197.49, "sgst": 197.49}',
                "portal_record": None,
                "match_reason": "Not in Portal: no portal note matched on reference, "
                                "or on GSTIN + date + value.",
            },
            {
                "run_id": run_id, "classification": "Matched",
                "confidence_score": 98.0, "confidence_band": "High",
                "note_reference": "SRI312", "note_reference_raw": "SRI-312/25-26",
                "note_kind": "Credit Note", "note_date": "2025-03-26",
                "note_tax": 1608.42, "note_value": 22583.0, "note_side": "both",
                "difference_type": None,
                "books_record": '{"gstin": "07BBBBB0000B1ZZ", "party_name": "Supplier B"}',
                "portal_record": '{"gstin": "07BBBBB0000B1ZZ", "party_name": "Supplier B"}',
                "match_reason": "Reference match: GSTIN 07BBBBB0000B1ZZ, note SRI312.",
            },
        ])
        conn.commit()
    finally:
        conn.close()
    return run_id


def test_the_report_note_card_uses_the_matched_status_from_the_pass(run_with_notes, db):
    from src.reconciliation import report_export as ex

    data = ex.build_export_data(run_with_notes, db_path=str(db))
    kpi = data["kpi"]
    assert kpi["note_source"] == "matched"
    assert kpi["note_itc_available"] is True
    assert kpi["credit_note_count"] == 2
    assert kpi["note_matched_count"] == 1

    # The card now carries a MATCHED STATUS per note, from the pass.
    by_reference = {n["reference"]: n for n in data["credit_notes"]}
    assert by_reference["CD-12/04-2025"]["classification"] == "Not in Portal"
    assert by_reference["SRI-312/25-26"]["classification"] == "Matched"
    assert by_reference["SRI-312/25-26"]["confidence"] == "High"

    # Note-level ITC: gross on every note, at-risk on the ones that did not match.
    assert kpi["note_itc_gross"] == pytest.approx(394.98 + 1608.42)
    assert kpi["note_itc_at_risk"] == pytest.approx(394.98)


def test_the_report_separates_note_itc_from_period_itc(run_with_notes, db):
    """The note figures must NOT leak into the invoice ITC figures — those read
    the invoice results, and this run has none."""
    from src.reconciliation import report_export as ex

    data = ex.build_export_data(run_with_notes, db_path=str(db))
    assert data["kpi"]["period_itc_total"] == 0.0
    assert data["kpi"]["invoice_count"] == 0
    assert data["invoices"] == []
    assert data["kpi"]["note_itc_gross"] > 0


def test_the_workbook_carries_the_note_status_and_the_note_itc_rows(run_with_notes, db, tmp_path, monkeypatch):
    from openpyxl import load_workbook

    from src.reconciliation import report_export as ex

    monkeypatch.setattr(ex, "_recalculate_workbook", lambda path: None)
    data = ex.build_export_data(run_with_notes, db_path=str(db))
    out = ex.write_xlsx(data, tmp_path / "notes_report.xlsx")

    wb = load_workbook(out, data_only=False)
    assert "Credit Notes" in wb.sheetnames
    sheet = wb["Credit Notes"]
    assert sheet.cell(row=1, column=15).value == "Classification"
    assert sheet.cell(row=2, column=15).value in ("Not in Portal", "Matched")
    # ...and column K is still the total-tax column the totals reference.
    assert sheet.cell(row=1, column=11).value == "Total tax"

    overview = wb["Overview"]
    labels = [str(overview.cell(row=r, column=1).value or "") for r in range(1, 30)]
    assert any(label.startswith("Note ITC (gross)") for label in labels)
    assert any(label.startswith("Note ITC at risk") for label in labels)
    # The invoice key figures did not move: the note block was APPENDED below
    # them, so B7 is still the ITC-at-stake SUMIF and B16 still counts the
    # Credit Notes sheet — the contract the verify scripts read.
    assert str(overview["A7"].value).startswith("ITC at stake")
    assert str(overview["B7"].value).startswith("=SUMIF(")
    assert str(overview["A16"].value) == "Credit notes"
    assert str(overview["B16"].value).startswith("=COUNTA('Credit Notes'")
