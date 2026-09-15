"""Shared helpers for Excel ingestion (Tier-A, no LLM).

Header detection is token-based and tolerant: column names vary between
portal exports ("GSTIN of Supplier", "Invoice Date") and Tally exports
("Voucher Number", "Party GSTIN"). We normalise header text and match
against per-field alias sets, so parsers survive cosmetic differences.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Iterable

import openpyxl


def _norm_header(text) -> str:
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text).strip().lower())


def _alias_match(alias: str, header: str) -> bool:
    """Word-boundary-aware alias match: 'gst' must not match inside 'GSTIN'."""
    if alias == header:
        return True
    pat = re.compile(rf"(^|[^a-z0-9]){re.escape(alias)}([^a-z0-9]|$)")
    return bool(pat.search(header))


def sheet_names(path: str | Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def load_workbook(path):
    return openpyxl.load_workbook(path, data_only=True)


def locate_header(
    ws,
    aliases: dict[str, tuple[str, ...]],
    scan_rows: int = 30,
    min_score: int = 2,
) -> tuple[int | None, dict[str, int]]:
    """Locate the header row by scoring token matches.

    Returns (header_row_idx, {field: col_index}) or (None, {}) when fewer
    than min_score fields match. The row with the most matched fields wins.
    """
    best: tuple[int, int, dict[str, int]] | None = None
    for r in range(1, min(scan_rows, (ws.max_row or 0) + 1)):
        col_map: dict[str, int] = {}
        for c in range(1, (ws.max_column or 0) + 1):
            h = _norm_header(ws.cell(row=r, column=c).value)
            if not h:
                continue
            for field, al in aliases.items():
                if field in col_map:
                    continue
                for a in al:
                    if _alias_match(a, h):
                        col_map[field] = c
                        break
        score = len(col_map)
        if score > 0 and (best is None or score > best[0]):
            best = (score, r, col_map)
    if best is None or best[0] < min_score:
        return None, {}
    return best[1], best[2]


def iter_rows(
    ws,
    header_row: int,
    col_map: dict[str, int],
    skip_total_rows: bool = True,
) -> Iterable[dict[str, Any]]:
    """Yield {field: raw_cell} dicts for tabular data below the header.

    Blank/merged rows are skipped; when skip_total_rows is set, rows whose
    first mapped column contains "total" / "grand total" are dropped too
    (common portal/Tally subtotal rows).
    """
    if not col_map:
        return
    key_col = next(iter(col_map.values()))
    for r in range(header_row + 1, (ws.max_row or 0) + 1):
        row = {f: ws.cell(row=r, column=c).value for f, c in col_map.items()}
        if not any(v is not None and str(v).strip() != "" for v in row.values()):
            continue
        if skip_total_rows:
            k = ws.cell(row=r, column=key_col).value
            if k is not None and re.search(r"total|subtotal|grand", str(k), re.I):
                continue
        yield row
