"""Excel export — one sheet per classification, plus a summary sheet.

Disposable PoC: no formatting effort beyond frozen header row, sensible
column widths, and wrapped `match_reason` text (the column an accountant
reads to decide whether the engine is right).
"""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.data_paths import DATA_ROOT
from src.queries import get_results, get_run, get_run_summary

_MATCH_REASON_WIDTH = 90
_DEFAULT_COL_WIDTH = 22
_WRAP_COLUMNS = {"match_reason"}


def _exports_dir(client: str, period: str) -> Path:
    d = DATA_ROOT / client / period / "exports"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _flatten_record_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Expand books_record / portal_record dicts into prefixed flat columns
    so an accountant can see the underlying fields directly in the sheet."""
    out = df.copy()
    for side in ("books_record", "portal_record"):
        if side not in out.columns:
            continue
        expanded = out[side].apply(lambda v: v if isinstance(v, dict) else {})
        expanded_df = pd.json_normalize(expanded).add_prefix(f"{side.split('_')[0]}_")
        expanded_df.index = out.index
        out = pd.concat([out.drop(columns=[side]), expanded_df], axis=1)
    if "matched_record_ids" in out.columns:
        out["matched_record_ids"] = out["matched_record_ids"].apply(
            lambda v: json.dumps(v) if isinstance(v, list) else v
        )
    return out


def _order_columns(df: pd.DataFrame) -> list[str]:
    """Put the accountant-relevant columns first, source fields after."""
    priority = [
        "result_id", "classification", "difference_type", "confidence_band",
        "confidence_score", "match_reason", "reviewed", "review_stale",
        "reviewer_note", "reviewed_at", "fingerprint",
    ]
    rest = [c for c in df.columns if c not in priority]
    ordered = [c for c in priority if c in df.columns] + rest
    return ordered


def _write_sheet(writer, sheet_name: str, df: pd.DataFrame) -> None:
    df = df.reset_index(drop=True)
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]

    # Freeze header row.
    ws.freeze_panes = "A2"

    # Bold header.
    for cell in ws[1]:
        cell.font = Font(bold=True)

    # Column widths + wrap for match_reason.
    for idx, col_name in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        if col_name in _WRAP_COLUMNS:
            ws.column_dimensions[letter].width = _MATCH_REASON_WIDTH
            for row in range(2, len(df) + 2):
                ws.cell(row=row, column=idx).alignment = Alignment(wrap_text=True, vertical="top")
        else:
            max_len = max(
                [len(str(col_name))] + [len(str(v)) for v in df[col_name].astype(str).head(200)]
            )
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), _DEFAULT_COL_WIDTH)


def _write_summary_sheet(writer, run: dict[str, Any], summary: dict[str, Any]) -> None:
    ws = writer.book.create_sheet("Summary", 0)

    ws["A1"] = "Run Summary"
    ws["A1"].font = Font(bold=True, size=14)

    rows = [
        ("Run ID", run["run_id"]),
        ("Client", run["client"]),
        ("Period", run["period"]),
        ("Recon Type", run["recon_type"]),
        ("Run Timestamp", run["run_timestamp"]),
        ("Source Files", ", ".join(run["source_file_names"])),
        ("Total Results", summary["total_results"]),
    ]
    r = 3
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Counts by classification").font = Font(bold=True)
    r += 1
    for cls, count in summary["by_classification"].items():
        value = summary["value_by_classification"].get(cls, 0.0)
        ws.cell(row=r, column=1, value=cls)
        ws.cell(row=r, column=2, value=count)
        ws.cell(row=r, column=3, value=f"₹{value:,.2f}")
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Counts by confidence band").font = Font(bold=True)
    r += 1
    for band, count in summary["by_confidence_band"].items():
        ws.cell(row=r, column=1, value=band)
        ws.cell(row=r, column=2, value=count)
        r += 1

    if summary["by_difference_type"]:
        r += 1
        ws.cell(row=r, column=1, value="Counts by difference_type").font = Font(bold=True)
        r += 1
        for dt, count in summary["by_difference_type"].items():
            ws.cell(row=r, column=1, value=dt)
            ws.cell(row=r, column=2, value=count)
            r += 1

    r += 2
    ws.cell(row=r, column=1, value="Config snapshot (as used for this run)").font = Font(bold=True)
    r += 1
    config_dict = yaml.safe_load(run["config_snapshot"])
    readable = yaml.dump(config_dict, default_flow_style=False, sort_keys=False)
    for line in readable.splitlines():
        ws.cell(row=r, column=1, value=line)
        r += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 20


def export_run(run_id: int, output_path: str | Path | None = None, *, db_path=None) -> Path:
    """Export a run to an Excel workbook: one sheet per classification plus
    a Summary sheet. Returns the path written.

    If output_path is None, writes to
    data/{client}/{period}/exports/{recon_type}_run{run_id}_{timestamp}.xlsx
    """
    run = get_run(run_id, db_path=db_path)
    if run is None:
        raise ValueError(f"Run {run_id} not found")

    results_df = get_results(run_id, db_path=db_path)
    summary = get_run_summary(run_id, db_path=db_path)

    if output_path is None:
        ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        fname = f"{run['recon_type']}_run{run_id}_{ts}.xlsx"
        output_path = _exports_dir(run["client"], run["period"]) / fname
    else:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

    flat_df = _flatten_record_columns(results_df)
    ordered_cols = _order_columns(flat_df)
    flat_df = flat_df[ordered_cols]

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for classification in sorted(flat_df["classification"].unique()):
            sheet_df = flat_df[flat_df["classification"] == classification]
            sheet_name = classification[:31]  # Excel sheet name limit
            _write_sheet(writer, sheet_name, sheet_df)

        _write_summary_sheet(writer, run, summary)

    return output_path
