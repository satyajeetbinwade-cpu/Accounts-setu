"""Export tab \u2014 generate an Excel workbook and offer it as a download.

Calls src.export.export_run unchanged (which writes the file to disk), then
reads the resulting bytes and surfaces a download button so staff don't
have to go find the file in the filesystem.
"""

from __future__ import annotations

from pathlib import Path

import streamlit as st
from openpyxl import load_workbook

from src import queries
from src.export import export_run
from src.ui import state


def render_export_tab() -> None:
    run_id = state.get_selected_run_id()
    if run_id is None:
        st.info(
            "Select a run in the sidebar before exporting. "
            "If no runs exist yet, go to the **Run** tab and execute one first."
        )
        return

    client = st.session_state.get("ctx_client")
    period = st.session_state.get("ctx_period")

    st.subheader(f"Export run {run_id}")
    st.write(
        "Generates an Excel workbook you can download, one sheet per classification "
        "plus a summary sheet. A copy is also kept on disk for the record."
    )

    st.markdown('<hr class="setu-section-rule" />', unsafe_allow_html=True)

    if st.button("Generate export", type="primary"):
        try:
            path = export_run(run_id)
        except Exception as exc:  # noqa: BLE001
            st.error(
                "The export couldn't be written. "
                "Check that the destination folder exists and is writable."
            )
            with st.expander("Details"):
                st.code(str(exc))
            return

        _render_download(path, run_id)


def _render_download(path: Path, run_id: int) -> None:
    """Confirm the file is complete, then offer it via a download button."""
    try:
        data = path.read_bytes()
        wb = load_workbook(path, read_only=True)
        sheet_count = len(wb.sheetnames)
        wb.close()
        size_kb = len(data) / 1024
    except Exception as exc:  # noqa: BLE001
        st.error(
            "The export was written but could not be read back for download. "
            "Check the file on disk."
        )
        with st.expander("Details"):
            st.code(str(exc))
        return

    run = queries.get_run(run_id)
    ts = (run or {}).get("run_timestamp", "").replace("-", "").replace(":", "").split(".")[0]
    recon_type = (run or {}).get("recon_type", "RECON")
    client = (run or {}).get("client", "client")
    period = (run or {}).get("period", "period")
    download_name = f"{client}_{period}_{recon_type}_run{run_id}_{ts}.xlsx"

    st.success(
        f"Export ready \u2014 {sheet_count} sheet(s), {size_kb:.1f} KB."
    )
    st.download_button(
        "Download Excel workbook",
        data=data,
        file_name=download_name,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )
    st.caption(f"Also saved to `{path}`.")
