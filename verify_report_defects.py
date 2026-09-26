"""§8 mandatory self-verification for the Run 11 report defect-fix build.

Run:

    venv/bin/python verify_report_defects.py

Prints one PASS/FAIL line per §8 item and exits non-zero if any fails.
It regenerates Run 11's three exports (HTML/PDF/Excel) from the fixed code,
then inspects them directly — the same way the defects were found.

The §2 cross-client check needs a SECOND client/dataset. It is created
idempotently here (client record + a copied dataset with a different supplier
GSTIN + one run) if it does not already exist.
"""

from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import fitz  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from src import db, queries, runner  # noqa: E402
from src.clients import service as clients  # noqa: E402
from src.config_loader import load_config  # noqa: E402
from src.reconciliation import report_export as ex  # noqa: E402
from src.reconciliation import review as rv  # noqa: E402
from src.reconciliation import run_model  # noqa: E402

RUN_ID = 11
CLIENT_GSTIN = "07AAACP9472A1ZJ"
SUPPLIER_GSTIN = "07AAXFS9006M1ZC"
GROSS = 12757.26
RESULTS: list[tuple[str, bool, str]] = []


def check(label: str, ok: bool, detail: str = "") -> None:
    RESULTS.append((label, ok, detail))
    print(f"{'PASS' if ok else 'FAIL'}  {label}" + (f"  — {detail}" if detail else ""))


def _donut_dashes(html: str, aria: str) -> tuple[list[float], float]:
    m = re.search(r'<svg[^>]*aria-label="' + re.escape(aria) + r'"[^>]*>(.*?)</svg>', html, re.S)
    if not m:
        return [], 0.0
    return [float(x) for x in re.findall(r'stroke-dasharray="([0-9.]+) ', m.group(1))], 2 * 3.141592653589793 * 70.0


def _approx(value, target: float, tol: float = 0.01) -> bool:
    try:
        return abs(float(value) - target) <= tol
    except (TypeError, ValueError):
        return False


def _ensure_second_dataset() -> str | None:
    """Create (idempotently) a second client with its own GSTIN and a copied
    dataset carrying a DIFFERENT supplier GSTIN; run it. Returns the folder."""
    folder = "Meridian Fabrics"
    own_gstin = "27AABCM1234K1Z9"
    existing = {c["legal_name"]: c["client_id"] for c in clients.list_clients()}
    if folder not in existing:
        clients.create_client(
            legal_name=folder, pan="AABCM1234K", assigned_team=None,
            initial_gstin=own_gstin, initial_state="Maharashtra", actor="admin",
        )
    cid = {c["legal_name"]: c["client_id"] for c in clients.list_clients()}[folder]

    src_dir = Path("data/Test Client/2026-08")
    dst_dir = Path("data") / folder / "2026-08"
    if not dst_dir.exists():
        import openpyxl

        def rewrite(src: Path, dst: Path) -> None:
            dst.parent.mkdir(parents=True, exist_ok=True)
            wb = openpyxl.load_workbook(src)
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and SUPPLIER_GSTIN in cell.value:
                            cell.value = cell.value.replace(SUPPLIER_GSTIN, "27AAXFS9006M1Z8")
            wb.save(dst)

        rewrite(src_dir / "gstr2b" / "082026_07AAACP9472A1ZJ_GSTR2B_15092026 (1).xlsx",
                dst_dir / "gstr2b" / "082026_07AAACP9472A1ZJ_GSTR2B_15092026 (1).xlsx")
        rewrite(src_dir / "tally" / "PurchaseRegister(Bill-wise) (1).xlsx",
                dst_dir / "tally" / "PurchaseRegister(Bill-wise) (1).xlsx")

    has_run = any(
        (r.get("client") or "") == folder and (r.get("recon_type") or "") == "GST"
        for r in queries.list_runs().to_dict("records")
    )
    if not has_run:
        runner.execute_run(
            folder, "2026-08", "GST", None, load_config(),
            selected_files={
                "gstr2b": "082026_07AAACP9472A1ZJ_GSTR2B_15092026 (1).xlsx",
                "tally": "PurchaseRegister(Bill-wise) (1).xlsx",
            },
            client_id=cid, actor="admin",
        )
    return folder


def main() -> int:
    db.init_db()
    run, model = run_model.load_run_model(RUN_ID)
    k = model["kpi"]
    data = ex.build_export_data(RUN_ID)
    ek = data["kpi"]

    # The three exports for Run 11, regenerated from the fixed code.
    paths = {fmt: ex.export_report(RUN_ID, fmt) for fmt in ("html", "pdf", "xlsx")}
    html = Path(paths["html"]).read_text(encoding="utf-8")
    pdf = fitz.open(paths["pdf"])
    pdf_text = "\n".join(p.get_text() for p in pdf)
    ws = load_workbook(paths["xlsx"], data_only=True)["Overview"]

    # --- 1. Classification donut: Matched is the large majority -------------
    dashes, circ = _donut_dashes(html, "Records")
    if len(dashes) >= 2 and circ:
        frac = dashes[0] / circ
        check(
            "1. Classification donut: Matched occupies the large majority",
            abs(frac - 35 / 39) < 0.02,
            f"Matched arc {frac * 100:.1f}% of the ring (expected 35/39 = 89.7%), "
            f"Not in Books {(1 - frac) * 100:.1f}%",
        )
    else:
        check("1. Classification donut: Matched occupies the large majority", False,
              "could not read the classification donut arcs from the HTML")

    # --- 2. Header GSTIN = the client's own, across all three formats -------
    client_gstin = run_model.client_gstin_for(model.get("client_id"))
    supplier_gstins = {i["gstin"] for i in data["invoices"] if i["gstin"]}
    # Check the HEADER field specifically: a supplier GSTIN legitimately
    # appears in the detail table, so "absent from the document" is not the
    # test — the header must carry the client's own GSTIN.
    html_header = re.findall(r"GSTIN ([0-9A-Z]{15})", html)
    pdf_header = re.findall(r"GSTIN ([0-9A-Z]{15})", pdf_text)
    xlsx_header = str(ws["A2"].value or "")
    check(
        "2. Header GSTIN is the client's own (07AAACP9472A1ZJ) in HTML, PDF and XLSX",
        client_gstin == CLIENT_GSTIN
        and html_header[:1] == [CLIENT_GSTIN]
        and pdf_header[:1] == [CLIENT_GSTIN]
        and CLIENT_GSTIN in xlsx_header
        and CLIENT_GSTIN not in supplier_gstins,
        f"client record={client_gstin!r}; header HTML={html_header[:1]}, "
        f"PDF={pdf_header[:1]}, XLSX={xlsx_header!r}; header never a supplier GSTIN",
    )

    # --- 3. Second client/dataset: header changes, never a supplier ---------
    try:
        second = _ensure_second_dataset()
        second_runs = [r for r in queries.list_runs().to_dict("records")
                       if (r.get("client") or "") == second and (r.get("recon_type") or "") == "GST"]
        second_run_id = int(max(r["run_id"] for r in second_runs))
        sdata = ex.build_export_data(second_run_id)
        s_suppliers = {i["gstin"] for i in sdata["invoices"] if i["gstin"]}
        check(
            "3. Second client/dataset: header GSTIN changes and never coincides with a supplier",
            sdata["gstin"] != CLIENT_GSTIN and bool(sdata["gstin"])
            and sdata["gstin"] not in s_suppliers,
            f"run {second_run_id} ({second}) header GSTIN={sdata['gstin']!r}, "
            f"suppliers={sorted(s_suppliers)[:4]}…",
        )
    except Exception as exc:  # noqa: BLE001
        check("3. Second client/dataset: header GSTIN changes", False, f"setup failed: {exc}")

    # --- 4. Excel cached values: every KPI cell a real number ---------------
    kpi_rows = {}
    for r in range(1, 40):
        label = str(ws.cell(row=r, column=1).value or "")
        if label and ws.cell(row=r, column=2).value is not None:
            kpi_rows[label] = ws.cell(row=r, column=2).value
    required = [
        "ITC at stake (tax on exceptions)", "Period ITC (total tax in run)",
        "ITC at stake % of period ITC", "Invoices", "Matched", "Exceptions",
        "Not in books", "Not in portal", "Amount difference", "Credit notes",
        "Credit note tax", "Gross invoice value on exceptions (context only)",
    ]
    blanks = [lbl for lbl in required if kpi_rows.get(lbl) in (None, "")]
    check(
        "4. Excel ships cached values — every KPI cell is a real number",
        not blanks,
        "all KPI cells populated by the recalculation pass" if not blanks
        else f"blank cells: {blanks}",
    )

    # --- 5. Gross value cell == 12,757.26 and matches PDF/HTML --------------
    gross_cell = kpi_rows.get("Gross invoice value on exceptions (context only)")
    gross_ok = _approx(gross_cell, GROSS)
    check(
        "5. Excel 'Gross invoice value on exceptions' == ₹12,757.26 and matches PDF/HTML",
        gross_ok and "12,757.26" in html and "12,757.26" in pdf_text,
        f"xlsx={gross_cell!r} · HTML={'12,757.26' in html} · PDF={'12,757.26' in pdf_text}",
    )

    # --- 6. IMS note Why/How distinct from the credit-note note -------------
    notes = {n.key: n for n in rv.data_quality_notes(
        model["items"], recon_type=model["recon_type"],
        source_files=run.get("source_file_names") or [],
        caveats=run.get("caveats") or [], run_notes=run.get("run_notes") or [],
    )}
    pairs = [(n.why, n.how) for n in notes.values()]
    ims = next((n for n in notes.values() if "IMS" in n.title), None)
    cn = next((n for n in notes.values() if "credit note" in n.title), None)
    check(
        "6. IMS note Why/How is distinct from the credit-note note's",
        ims is not None and cn is not None and (ims.why, ims.how) != (cn.why, cn.how)
        and len(pairs) == len(set(pairs)),
        f"IMS why reads {ims.why[:48]!r}…" if ims else "IMS note missing",
    )

    # --- 7. No regression to items already fixed ---------------------------
    itc = model["itc"] or {}
    itc_arcs = _donut_dashes(html, (itc.get("single_label") or "Total ITC"))[0]
    rounding_note = any("Rounding adjustments" in n.title for n in notes.values())
    confidence_scored = all(
        i["confidence_score"] for i in model["items"] if i["classification"] == "Matched"
    )
    check(
        "7. No regression: ITC-at-stake headline+%, ITC donut single ring, "
        "credit notes 12 vs 39, match-confidence, rounding note",
        k["itc_at_stake_display"] in html and k["itc_at_stake_pct_display"].endswith("%")
        and bool(itc.get("single")) and len(itc.get("slices") or []) == 1
        and not itc_arcs  # a single-colour ring draws one circle, no arcs
        and ek["credit_note_count"] == 12 and k["total_count"] == 39
        and confidence_scored and rounding_note,
        f"headline {k['itc_at_stake_display']} ({k['itc_at_stake_pct_display']}) · "
        f"ITC slices={[s['name'] for s in itc.get('slices') or []]} "
        f"(single ring: {'no arcs' if not itc_arcs else 'ARCS PRESENT'}) · "
        f"{k['total_count']} invoices vs {ek['credit_note_count']} credit notes · "
        f"confidence={'ok' if confidence_scored else 'MISSING'} · "
        f"rounding note={'present' if rounding_note else 'MISSING'}",
    )

    print()
    failed = [r for r in RESULTS if not r[1]]
    print(f"===== {len(RESULTS) - len(failed)}/{len(RESULTS)} PASS =====")
    for label, _, _ in failed:
        print(f"  FAILED: {label}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
