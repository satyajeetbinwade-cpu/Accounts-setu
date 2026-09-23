"""Phase-1 tools state — Run / Review / Compare / Config / Export.

Every handler calls ``src.*`` services (runner / queries / export /
config_loader / discovery / ingestion_ai); no business logic lives here.
"""

from __future__ import annotations

import base64
import difflib
from dataclasses import dataclass

import reflex as rx

from src import queries
from src.auth import service as auth
from src.config_loader import CONFIG_PATH, load_config
from src.export import export_run
from src.f5 import service as f5
from src.ingestion_ai import service as ingestion_ai
from src.runner import RunExecutionError, execute_run
from src.shared import discovery
from setu.state.auth_state import AuthState

AT_RISK_CLASSIFICATIONS = ("Amount Difference", "Not in Books", "Not in Portal")
CLASSIFICATION_ORDER = ["Amount Difference", "Not in Portal", "Not in Books", "Matched"]


@dataclass
class ContextOption:
    value: str
    label: str


@dataclass
class SourceSlot:
    source_type: str
    label: str
    hint: str
    is_books: bool
    files: list[str]
    selected: str
    present: bool
    runnable: bool
    reason: str


@dataclass
class RunSummary:
    run_id: int
    client: str
    period: str
    recon_type: str
    timestamp: str
    total: int


@dataclass
class ClassCount:
    label: str
    count: int
    value: str


@dataclass
class ResultRow:
    result_id: int
    fingerprint: str
    classification: str
    difference_type: str
    confidence_band: str
    confidence_score: int
    review_state: str
    match_reason: str
    books_identity: str
    portal_identity: str
    value: str


@dataclass
class RecordField:
    label: str
    value: str
    differs: bool


@dataclass
class ChangeRow:
    classification_a: str
    difference_type_a: str
    classification_b: str
    difference_type_b: str
    match_reason_a: str
    match_reason_b: str


@dataclass
class MovementRow:
    from_label: str
    to_label: str
    count: int
    moved: bool


class Phase1State(AuthState):
    """Shared state for the Phase-1 tools."""

    # ------------------------------------------------------------------
    # Context (client / period / recon type / selected run)
    # ------------------------------------------------------------------
    clients: list[str] = []
    periods: list[str] = []
    recon_types: list[str] = []
    runs: list[RunSummary] = []

    ctx_client: str = ""
    ctx_period: str = ""
    ctx_recon_type: str = ""
    selected_run_id: int = 0

    flash: str = ""
    error: str = ""

    # ------------------------------------------------------------------
    # Run tool
    # ------------------------------------------------------------------
    slots: list[SourceSlot] = []
    executing: bool = False
    run_result: str = ""
    run_caveats: list[str] = []

    # ------------------------------------------------------------------
    # Review tool
    # ------------------------------------------------------------------
    review_filter_classification: str = "All"
    review_filter_band: str = "All"
    review_filter_state: str = "All"
    review_rows: list[ResultRow] = []
    review_total: int = 0
    review_reviewed: int = 0
    review_unreviewed: int = 0
    review_stale: int = 0
    at_risk_value: str = ""
    class_counts: list[ClassCount] = []

    selected_result_id: int = 0
    detail_classification: str = ""
    detail_difference_type: str = ""
    detail_band: str = ""
    detail_score: int = 0
    detail_reason: str = ""
    detail_review_state: str = ""
    detail_books: list[RecordField] = []
    detail_portal: list[RecordField] = []
    review_note: str = ""

    # ------------------------------------------------------------------
    # Compare tool
    # ------------------------------------------------------------------
    compare_run_a: int = 0
    compare_run_b: int = 0
    compare_changed: int = 0
    compare_unchanged: int = 0
    compare_change_rate: str = ""
    compare_only_a: int = 0
    compare_only_b: int = 0
    compare_movement: list[MovementRow] = []
    compare_changes: list[ChangeRow] = []
    compare_ready: bool = False
    compare_error: str = ""

    # ------------------------------------------------------------------
    # Config tool
    # ------------------------------------------------------------------
    config_current: str = ""
    config_snapshot: str = ""
    config_same: bool = False
    config_diff: str = ""
    config_error: str = ""

    # ------------------------------------------------------------------
    # Export tool
    # ------------------------------------------------------------------
    export_b64: str = ""
    export_name: str = ""
    export_sheets: int = 0
    export_size_kb: str = ""

    # ==================================================================
    # Context management
    # ==================================================================
    @rx.event
    def load(self):
        self.clients = discovery.list_clients()
        if not self.ctx_client and self.clients:
            self.ctx_client = self.clients[0]
        self._load_context()

    def _load_context(self) -> None:
        if not self.ctx_client:
            return
        self.periods = discovery.list_periods(self.ctx_client)
        if self.ctx_period not in self.periods:
            self.ctx_period = self.periods[0] if self.periods else ""
        if self.ctx_client and self.ctx_period:
            self.recon_types = discovery.list_recon_types(self.ctx_client, self.ctx_period)
        else:
            self.recon_types = []
        if self.ctx_recon_type not in self.recon_types:
            self.ctx_recon_type = self.recon_types[0] if self.recon_types else ""
        self._load_runs()
        self._on_context_changed()

    def _on_context_changed(self) -> None:
        self._load_slots()
        # Default the selected run to the newest for this context BEFORE
        # loading anything that depends on it.
        ids = [r.run_id for r in self.runs]
        if ids and self.selected_run_id not in ids:
            self.selected_run_id = ids[0]
        if not ids:
            self.selected_run_id = 0
        self._load_review()
        self._load_config()
        self._load_compare()
        self._reset_export()

    @rx.event
    def set_ctx_client(self, v: str):
        self.ctx_client = v
        self.selected_run_id = 0
        self._load_context()

    @rx.event
    def set_ctx_period(self, v: str):
        self.ctx_period = v
        self.selected_run_id = 0
        self._load_context()

    @rx.event
    def set_ctx_recon_type(self, v: str):
        self.ctx_recon_type = v
        self.selected_run_id = 0
        self._load_runs()
        self._on_context_changed()

    def _load_runs(self) -> None:
        if not (self.ctx_client and self.ctx_period and self.ctx_recon_type):
            self.runs = []
            return
        df = queries.list_runs(
            client=self.ctx_client, period=self.ctx_period, recon_type=self.ctx_recon_type
        )
        rows: list[RunSummary] = []
        for _, r in df.iterrows():
            rows.append(
                RunSummary(
                    run_id=int(r["run_id"]),
                    client=r["client"],
                    period=r["period"],
                    recon_type=r["recon_type"],
                    timestamp=str(r.get("run_timestamp") or "").replace("T", " ").split(".")[0],
                    total=0,
                )
            )
        self.runs = rows

    @rx.event
    def set_selected_run(self, run_id: str):
        self.selected_run_id = _to_int(run_id)
        self._load_review()
        self._load_config()
        self._reset_export()

    @rx.var
    def has_context(self) -> bool:
        return bool(self.ctx_client and self.ctx_period and self.ctx_recon_type)

    @rx.var
    def run_ids(self) -> list[str]:
        return [str(r.run_id) for r in self.runs]

    # ==================================================================
    # Run tool
    # ==================================================================
    def _load_slots(self) -> None:
        if not (self.ctx_client and self.ctx_period and self.ctx_recon_type):
            self.slots = []
            return
        status = discovery.source_file_status(self.ctx_client, self.ctx_period, self.ctx_recon_type)
        books_types = set(discovery.BOOKS_SOURCE_TYPES)
        out: list[SourceSlot] = []
        for source_type in discovery.RECON_SOURCE_TYPES[self.ctx_recon_type]:
            info = status.get(source_type, {})
            files = list(info.get("files") or [])
            selected = files[0] if files else ""
            runnable, reason = True, ""
            if selected:
                runnable, reason = self._file_runnable(source_type, selected)
            out.append(
                SourceSlot(
                    source_type=source_type,
                    label=discovery.source_type_label(source_type),
                    hint=discovery.source_type_hint(source_type),
                    is_books=source_type in books_types,
                    files=files,
                    selected=selected,
                    present=bool(files),
                    runnable=runnable,
                    reason=reason or "",
                )
            )
        self.slots = out

    def _file_runnable(self, source_type: str, filename: str) -> tuple[bool, str]:
        try:
            stored = ingestion_ai.get_stored_result(
                self.ctx_client, self.ctx_period, source_type, filename
            )
        except Exception as exc:  # noqa: BLE001
            return False, f"Couldn't read ingestion state: {exc}"
        if stored is None:
            return False, (
                "This file hasn't been through AI ingestion yet. Open it in Smart Ingestion "
                "so its columns can be mapped."
            )
        status = stored.get("status")
        if status in ("wrong_slot", "unrecognized"):
            return False, stored.get("message") or "This file doesn't look like data for this slot."
        if status not in ("ok", "partial", "blocked"):
            return False, f"This file is in state {status!r} and can't be used yet."
        return True, ""

    @rx.event
    def select_slot_file(self, source_type: str, filename: str):
        slot = next((s for s in self.slots if s.source_type == source_type), None)
        if slot is None:
            return
        runnable, reason = True, ""
        if filename:
            runnable, reason = self._file_runnable(source_type, filename)
        self.slots = [
            _replace_slot(
                s, filename=filename,
                runnable=runnable if s.source_type == source_type else s.runnable,
                reason=reason if s.source_type == source_type else s.reason,
            )
            if s.source_type == source_type else s
            for s in self.slots
        ]

    @rx.var
    def books_ready(self) -> bool:
        return any(s.is_books and s.selected for s in self.slots)

    @rx.var
    def portal_ready(self) -> bool:
        return any((not s.is_books) and s.selected for s in self.slots)

    @rx.var
    def run_blockers(self) -> list[str]:
        out: list[str] = []
        if not self.books_ready:
            out = out + ["Select a file for the books side."]
        if not self.portal_ready:
            out = out + ["Select a file for at least one portal source."]
        for s in self.slots:
            if s.selected and not s.runnable:
                out = out + [f"{s.label} — {s.reason}"]
        return out

    @rx.event
    def execute_run(self):
        self.error = ""
        self.flash = ""
        self.run_result = ""
        self.run_caveats = []
        if not (self.ctx_client and self.ctx_period and self.ctx_recon_type):
            self.error = "Pick a client, period, and recon type first."
            return
        selected_files = {s.source_type: s.selected for s in self.slots if s.selected}
        client_id = _client_id_for_folder(self.ctx_client)
        try:
            config = load_config()
            run_id = execute_run(
                self.ctx_client, self.ctx_period, self.ctx_recon_type, None, config,
                selected_files=selected_files, client_id=client_id, actor=self.username,
            )
        except RunExecutionError as exc:
            self.error = str(exc)
            return
        except Exception as exc:  # noqa: BLE001
            self.error = str(exc)
            return

        # F5 sync-health retrofit (best-effort).
        _record_sync_health(client_id, selected_files, succeeded=True, actor=self.username)
        # Module 2 retrofit — materialize the exception queue (best-effort).
        if client_id is not None:
            try:
                from src.module2 import service as m2

                m2.generate_exceptions_for_run(run_id, client_id=client_id, actor=self.username)
            except Exception:  # noqa: BLE001
                pass

        run = queries.get_run(run_id) or {}
        self.run_caveats = [c.get("label") or c.get("code") or "" for c in (run.get("caveats") or [])]
        self.flash = f"Run {run_id} complete for {self.ctx_client} / {self.ctx_period} / {self.ctx_recon_type}."
        self.run_result = str(run_id)
        self._load_runs()
        self.selected_run_id = run_id
        self._load_review()

    # ==================================================================
    # Review tool
    # ==================================================================
    def _load_review(self) -> None:
        if not self.selected_run_id:
            self.review_rows = []
            self.review_total = 0
            self.class_counts = []
            return
        df = queries.get_results(self.selected_run_id)
        summary = queries.get_run_summary(self.selected_run_id)

        value_by_cls = summary.get("value_by_classification", {})
        self.class_counts = [
            ClassCount(label=cls, count=int(summary["by_classification"].get(cls, 0)), value=_money(value_by_cls.get(cls, 0.0)))
            for cls in CLASSIFICATION_ORDER
        ]
        self.at_risk_value = _money(
            sum(value_by_cls.get(cls, 0.0) for cls in AT_RISK_CLASSIFICATIONS)
        )
        self.review_total = int(summary.get("total_results", 0))

        states: list[str] = []
        for _, row in df.iterrows():
            if bool(row.get("review_stale")):
                states.append("stale")
            elif bool(row.get("reviewed")):
                states.append("reviewed")
            else:
                states.append("unreviewed")
        self.review_reviewed = states.count("reviewed")
        self.review_unreviewed = states.count("unreviewed")
        self.review_stale = states.count("stale")

        rows: list[ResultRow] = []
        for (_, row), state_name in zip(df.iterrows(), states):
            if self.review_filter_classification != "All" and row["classification"] != self.review_filter_classification:
                continue
            if self.review_filter_band != "All" and row["confidence_band"] != self.review_filter_band:
                continue
            if self.review_filter_state != "All" and state_name != self.review_filter_state:
                continue
            books = row.get("books_record")
            portal = row.get("portal_record")
            rows.append(
                ResultRow(
                    result_id=int(row["result_id"]),
                    fingerprint=row.get("fingerprint") or "",
                    classification=row["classification"],
                    difference_type=row.get("difference_type") if isinstance(row.get("difference_type"), str) else "",
                    confidence_band=row["confidence_band"],
                    confidence_score=int(row["confidence_score"] or 0),
                    review_state=state_name,
                    match_reason=row.get("match_reason") or "",
                    books_identity=_identity_line(books, self.ctx_recon_type),
                    portal_identity=_identity_line(portal, self.ctx_recon_type),
                    value=_money(_row_value(books, portal)),
                )
            )
        self.review_rows = rows
        if self.selected_result_id:
            self._load_detail()

    @rx.event
    def set_review_filter_classification(self, v: str):
        self.review_filter_classification = v
        self.selected_result_id = 0
        self._load_review()

    @rx.event
    def set_review_filter_band(self, v: str):
        self.review_filter_band = v
        self._load_review()

    @rx.event
    def set_review_filter_state(self, v: str):
        self.review_filter_state = v
        self._load_review()

    @rx.event
    def open_result(self, result_id: int):
        self.selected_result_id = result_id
        self.review_note = ""
        self._load_detail()

    @rx.event
    def close_result(self):
        self.selected_result_id = 0

    def _load_detail(self) -> None:
        df = queries.get_results(self.selected_run_id)
        match = df[df["result_id"] == self.selected_result_id]
        if match.empty:
            self.selected_result_id = 0
            return
        row = match.iloc[0]
        self.detail_classification = row["classification"]
        self.detail_difference_type = row["difference_type"] if isinstance(row.get("difference_type"), str) else ""
        self.detail_band = row["confidence_band"]
        self.detail_score = int(row["confidence_score"] or 0)
        self.detail_reason = row.get("match_reason") or ""
        if bool(row.get("review_stale")):
            self.detail_review_state = "stale"
        elif bool(row.get("reviewed")):
            self.detail_review_state = "reviewed"
        else:
            self.detail_review_state = "unreviewed"
        self.review_note = row.get("reviewer_note") if isinstance(row.get("reviewer_note"), str) else ""
        books = row.get("books_record")
        portal = row.get("portal_record")
        differing = _diff_fields(books, portal)
        self.detail_books = _record_fields(books, differing)
        self.detail_portal = _record_fields(portal, differing)

    def set_review_note(self, v: str):
        self.review_note = v

    @rx.event
    def mark_reviewed(self):
        row = self._detail_row()
        if row is None:
            return
        queries.mark_reviewed(
            self.ctx_client, self.ctx_period, self.ctx_recon_type,
            row.get("fingerprint"), note=self.review_note or None, result=row,
        )
        self.flash = "Marked reviewed."
        self.selected_result_id = 0
        self._load_review()

    @rx.event
    def clear_review(self):
        row = self._detail_row()
        if row is None:
            return
        queries.clear_review(
            self.ctx_client, self.ctx_period, self.ctx_recon_type, row.get("fingerprint")
        )
        self.flash = "Review cleared."
        self.selected_result_id = 0
        self._load_review()

    def _detail_row(self) -> dict | None:
        df = queries.get_results(self.selected_run_id)
        match = df[df["result_id"] == self.selected_result_id]
        return None if match.empty else match.iloc[0].to_dict()

    # ==================================================================
    # Compare tool
    # ==================================================================
    def _load_compare(self) -> None:
        self.compare_error = ""
        ids = [r.run_id for r in self.runs]
        self.compare_ready = len(ids) >= 2
        if not self.compare_ready:
            self.compare_changed = self.compare_unchanged = 0
            self.compare_movement = []
            self.compare_changes = []
            return
        if self.compare_run_a not in ids:
            self.compare_run_a = ids[1] if len(ids) > 1 else ids[0]
        if self.compare_run_b not in ids:
            self.compare_run_b = ids[0]
        if self.compare_run_a == self.compare_run_b:
            self.compare_changes = []
            return
        try:
            result = queries.compare_runs(self.compare_run_a, self.compare_run_b)
        except Exception as exc:  # noqa: BLE001
            self.compare_error = str(exc)
            return
        self.compare_changed = result["changed_count"]
        self.compare_unchanged = result["unchanged_count"]
        total = max(self.compare_changed + self.compare_unchanged, 1)
        self.compare_change_rate = f"{self.compare_changed / total * 100:.1f}%"
        self.compare_only_a = len(result["only_in_a"])
        self.compare_only_b = len(result["only_in_b"])

        movement: list[MovementRow] = []
        for from_cls, tos in result["movement"].items():
            for to_cls, count in tos.items():
                movement.append(
                    MovementRow(from_label=from_cls, to_label=to_cls, count=int(count), moved=from_cls != to_cls)
                )
        movement.sort(key=lambda m: (not m.moved, -m.count))
        self.compare_movement = movement

        changes: list[ChangeRow] = []
        for _, row in result["changed"].head(60).iterrows():
            changes.append(
                ChangeRow(
                    classification_a=row["classification_a"],
                    difference_type_a=row["difference_type_a"] if isinstance(row.get("difference_type_a"), str) else "",
                    classification_b=row["classification_b"],
                    difference_type_b=row["difference_type_b"] if isinstance(row.get("difference_type_b"), str) else "",
                    match_reason_a=row.get("match_reason_a") or "",
                    match_reason_b=row.get("match_reason_b") or "",
                )
            )
        self.compare_changes = changes

    @rx.event
    def set_compare_a(self, v: str):
        self.compare_run_a = _to_int(v)
        self._load_compare()

    @rx.event
    def set_compare_b(self, v: str):
        self.compare_run_b = _to_int(v)
        self._load_compare()

    # ==================================================================
    # Config tool
    # ==================================================================
    def _load_config(self) -> None:
        self.config_error = ""
        try:
            self.config_current = CONFIG_PATH.read_text(encoding="utf-8")
        except Exception as exc:  # noqa: BLE001
            self.config_error = str(exc)
            self.config_current = ""
            return
        if not self.selected_run_id:
            self.config_snapshot = ""
            self.config_same = False
            self.config_diff = ""
            return
        run = queries.get_run(self.selected_run_id) or {}
        self.config_snapshot = run.get("config_snapshot") or ""
        self.config_same = self.config_current == self.config_snapshot
        if self.config_same:
            self.config_diff = ""
        else:
            lines = list(
                difflib.unified_diff(
                    self.config_snapshot.splitlines(),
                    self.config_current.splitlines(),
                    fromfile=f"run_{self.selected_run_id}_snapshot",
                    tofile="current",
                    lineterm="",
                )
            )
            self.config_diff = "\n".join(lines) or "(no textual diff)"

    # ==================================================================
    # Export tool
    # ==================================================================
    def _reset_export(self) -> None:
        self.export_b64 = ""
        self.export_name = ""
        self.export_sheets = 0
        self.export_size_kb = ""

    @rx.event
    def generate_export(self):
        self.error = ""
        self.flash = ""
        if not self.selected_run_id:
            self.error = "Select a run to export."
            return
        try:
            path = export_run(self.selected_run_id)
            data = path.read_bytes()
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            self.export_sheets = len(wb.sheetnames)
            wb.close()
        except Exception as exc:  # noqa: BLE001
            self.error = f"The export couldn't be written: {exc}"
            return
        run = queries.get_run(self.selected_run_id) or {}
        ts = str(run.get("run_timestamp") or "").replace("-", "").replace(":", "").split(".")[0]
        self.export_name = (
            f"{run.get('client', 'client')}_{run.get('period', 'period')}_"
            f"{run.get('recon_type', 'RECON')}_run{self.selected_run_id}_{ts}.xlsx"
        )
        self.export_b64 = base64.b64encode(data).decode()
        self.export_size_kb = f"{len(data) / 1024:.1f}"
        self.flash = f"Export ready — {self.export_sheets} sheet(s), {self.export_size_kb} KB."


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _money(v) -> str:
    try:
        return f"₹{float(v):,.2f}"
    except (TypeError, ValueError):
        return "—"


def _to_int(v) -> int:
    try:
        return int(v)
    except (TypeError, ValueError):
        return 0


def _replace_slot(s: SourceSlot, *, filename: str, **kw) -> SourceSlot:
    return SourceSlot(
        source_type=s.source_type,
        label=s.label,
        hint=s.hint,
        is_books=s.is_books,
        files=s.files,
        selected=filename,
        present=s.present,
        runnable=kw.get("runnable", s.runnable),
        reason=kw.get("reason", s.reason),
    )


def _identity_line(record, recon_type: str) -> str:
    if not isinstance(record, dict):
        return "— no record —"
    if recon_type == "TDS":
        keys = ("deductee_name", "deductor_name", "pan", "section", "amount_paid_credited")
    elif recon_type == "OTHER":
        keys = ("reference", "date", "amount", "party_name")
    else:
        keys = ("invoice_number", "gstin", "party_name", "invoice_value")
    parts = [f"{k}={record[k]}" for k in keys if record.get(k) not in (None, "")]
    return " · ".join(parts) if parts else "—"


def _row_value(books, portal) -> float:
    for rec in (books, portal):
        if isinstance(rec, dict):
            for key in ("invoice_value", "amount_paid_credited", "tax_deducted", "amount"):
                if rec.get(key) is not None:
                    try:
                        return abs(float(rec[key]))
                    except (TypeError, ValueError):
                        continue
    return 0.0


def _diff_fields(books, portal) -> set[str]:
    if not isinstance(books, dict) or not isinstance(portal, dict):
        return set()
    keys = (set(books) | set(portal)) - {"original_row"}
    out = set()
    for k in keys:
        bv, pv = books.get(k), portal.get(k)
        bs, ps = ("" if bv is None else str(bv)), ("" if pv is None else str(pv))
        if bs != ps:
            out.add(k)
    return out


def _record_fields(record, differing: set[str]) -> list[RecordField]:
    if not isinstance(record, dict):
        return []
    out: list[RecordField] = []
    for k, v in record.items():
        if k == "original_row":
            continue
        out.append(
            RecordField(
                label=str(k).replace("_", " "),
                value="" if v is None else str(v),
                differs=k in differing,
            )
        )
    return out


def _client_id_for_folder(folder: str) -> int | None:
    from src.clients import service as clients

    def norm(name: str) -> str:
        return "".join(ch for ch in str(name).lower() if ch.isalnum())

    try:
        candidates = clients.list_clients(include_inactive=True)
    except Exception:  # noqa: BLE001
        return None
    target = norm(folder)
    for c in candidates:
        if norm(c["legal_name"]) == target:
            return c["client_id"]
    for c in candidates:
        n = norm(c["legal_name"])
        if target and (n.startswith(target) or target.startswith(n)):
            return c["client_id"]
    return None


_SYNC_SOURCE_FOR = {
    "tally": "tally", "gstr2b": "gst_portal", "ims": "gst_portal",
    "form26as": "traces", "tds": "traces",
    "bank": "bank", "vendor_ledger": "tally", "opening_balances": "tally",
    "loan_sheet": "tally", "salary": "tally",
}


def _record_sync_health(client_id, selected_files: dict, *, succeeded: bool, actor: str) -> None:
    if client_id is None:
        return
    try:
        for source_type in selected_files:
            src = _SYNC_SOURCE_FOR.get(source_type)
            if src is None:
                continue
            f5.record_sync_health(
                client_id=client_id, source=src,
                status="succeeded" if succeeded else "attempted_failed",
            )
    except Exception:  # noqa: BLE001
        pass
